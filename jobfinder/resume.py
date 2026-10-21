from __future__ import annotations

from pathlib import Path


def read_resume_text(path: str) -> str:
    """Read a resume (.docx, .pdf, .txt, .md) as plain text. Nothing leaves this computer."""
    file = Path(path)
    suffix = file.suffix.lower()
    if suffix == ".docx":
        try:
            import docx  # python-docx
        except ImportError:
            raise SystemExit("Reading .docx needs python-docx. Run: pip install python-docx")
        document = docx.Document(str(file))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    if suffix == ".pdf":
        try:
            import pypdf
        except ImportError:
            raise SystemExit("Reading .pdf needs pypdf. Run: pip install pypdf")
        reader = pypdf.PdfReader(str(file))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix in {".txt", ".md"}:
        return file.read_text(encoding="utf-8", errors="ignore")
    raise SystemExit(f"Unsupported resume type '{suffix}'. Use .docx, .pdf, .txt or .md")
